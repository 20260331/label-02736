# -*- coding: utf-8 -*-
"""
Excel 存储类 - 数据持久化层
"""
import os
import threading
from openpyxl import Workbook, load_workbook
from app.utils.logger import get_logger

logger = get_logger(__name__)


class ExcelStorage:
    """Excel 文件存储类"""
    
    _lock = threading.Lock()
    
    def __init__(self, file_path, columns):
        """
        初始化 Excel 存储
        
        Args:
            file_path: Excel 文件路径
            columns: 列名列表
        """
        self.file_path = file_path
        self.columns = columns
        self._ensure_file_exists()
    
    def _ensure_file_exists(self):
        """确保文件存在，不存在则创建"""
        with self._lock:
            if not os.path.exists(self.file_path):
                # 确保目录存在
                os.makedirs(os.path.dirname(self.file_path), exist_ok=True)
                
                # 创建新的工作簿
                wb = Workbook()
                ws = wb.active
                ws.title = 'data'
                
                # 写入表头
                for col_idx, col_name in enumerate(self.columns, 1):
                    ws.cell(row=1, column=col_idx, value=col_name)
                
                wb.save(self.file_path)
                logger.info(f'Created new Excel file: {self.file_path}')
    
    def read_all(self):
        """读取所有数据"""
        with self._lock:
            try:
                wb = load_workbook(self.file_path)
                ws = wb.active
                
                data = []
                # 跳过表头，从第2行开始读取
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[0] is None:  # 跳过空行
                        continue
                    row_data = {}
                    for col_idx, col_name in enumerate(self.columns):
                        value = row[col_idx] if col_idx < len(row) else None
                        row_data[col_name] = value
                    data.append(row_data)
                
                wb.close()
                return data
            except Exception as e:
                logger.error(f'Error reading Excel file: {e}')
                raise
    
    def find_by_id(self, id_value, id_column='id'):
        """根据 ID 查找单条记录"""
        all_data = self.read_all()
        for item in all_data:
            if str(item.get(id_column)) == str(id_value):
                return item
        return None
    
    def find_by_field(self, field_name, field_value):
        """根据字段值查找记录"""
        all_data = self.read_all()
        results = []
        for item in all_data:
            if str(item.get(field_name)) == str(field_value):
                results.append(item)
        return results
    
    def insert(self, data):
        """插入一条记录"""
        with self._lock:
            try:
                wb = load_workbook(self.file_path)
                ws = wb.active
                
                # 找到下一个空行
                next_row = ws.max_row + 1
                
                # 写入数据
                for col_idx, col_name in enumerate(self.columns, 1):
                    value = data.get(col_name)
                    ws.cell(row=next_row, column=col_idx, value=value)
                
                wb.save(self.file_path)
                wb.close()
                logger.info(f'Inserted record: {data.get("id")}')
                return True
            except Exception as e:
                logger.error(f'Error inserting record: {e}')
                raise
    
    def update(self, id_value, data, id_column='id'):
        """更新一条记录"""
        with self._lock:
            try:
                wb = load_workbook(self.file_path)
                ws = wb.active
                
                # 找到 ID 列的索引
                id_col_idx = self.columns.index(id_column) + 1
                
                # 查找并更新记录
                found = False
                for row_idx in range(2, ws.max_row + 1):
                    cell_value = ws.cell(row=row_idx, column=id_col_idx).value
                    if str(cell_value) == str(id_value):
                        # 更新数据
                        for col_idx, col_name in enumerate(self.columns, 1):
                            if col_name in data:
                                ws.cell(row=row_idx, column=col_idx, value=data[col_name])
                        found = True
                        break
                
                if found:
                    wb.save(self.file_path)
                    logger.info(f'Updated record: {id_value}')
                
                wb.close()
                return found
            except Exception as e:
                logger.error(f'Error updating record: {e}')
                raise
    
    def delete(self, id_value, id_column='id'):
        """删除一条记录"""
        with self._lock:
            try:
                wb = load_workbook(self.file_path)
                ws = wb.active
                
                # 找到 ID 列的索引
                id_col_idx = self.columns.index(id_column) + 1
                
                # 查找并删除记录
                found = False
                for row_idx in range(2, ws.max_row + 1):
                    cell_value = ws.cell(row=row_idx, column=id_col_idx).value
                    if str(cell_value) == str(id_value):
                        ws.delete_rows(row_idx)
                        found = True
                        break
                
                if found:
                    wb.save(self.file_path)
                    logger.info(f'Deleted record: {id_value}')
                
                wb.close()
                return found
            except Exception as e:
                logger.error(f'Error deleting record: {e}')
                raise
    
    def search(self, keyword, search_columns):
        """搜索记录"""
        all_data = self.read_all()
        if not keyword:
            return all_data
        
        keyword_lower = str(keyword).lower()
        results = []
        for item in all_data:
            for col in search_columns:
                value = item.get(col)
                if value and keyword_lower in str(value).lower():
                    results.append(item)
                    break
        return results
    
    def count(self):
        """获取记录总数"""
        return len(self.read_all())
